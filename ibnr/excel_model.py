"""Live-formula Excel version of the valuation.

Python remains the engine of record. This workbook is what a valuation team
hands around: the triangle is data, and every link ratio, average,
completion factor, ultimate, IBNR and load is an Excel formula. Typing a
link ratio into a yellow override cell re-runs the reserve. The Summary
sheet carries the Python result next to the workbook's, so any divergence
is visible.

Conventions: blue font = input data, yellow fill = judgment input,
black = formula, grey fill = link ratio excluded from the average.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col

from .valuation import LARGE, CategoryResult, Valuation

HEADER = PatternFill("solid", fgColor="256ABF")
INPUT = Font(color="1C5CAB")
JUDGMENT = PatternFill("solid", fgColor="FFF2B3")
EXCLUDED = PatternFill("solid", fgColor="E1E0D9")
BOLD = Font(bold=True)
ABBREV = {"Inpatient": "IP", "Outpatient": "OP", "Professional": "Prof", LARGE: "LC"}


def _header(ws, labels: list[str], row: int = 1) -> None:
    for j, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def _ref(sheet: str, r: int, c: int) -> str:
    return f"'{sheet}'!{col(c)}{r}"


def _assumptions(wb: Workbook, v: Valuation) -> dict[str, str]:
    ws = wb.create_sheet("Assumptions")
    cfg = v.config
    _header(ws, ["Assumption", "Value", "Note"])
    rows = [
        ("Valuation month", str(v.valuation_month), "Paid-through month"),
        ("PMPM / B-F months", cfg.pmpm_lags, "Incurred months younger than this lag use B-F on the PMPM expected"),
        ("Margin %", cfg.margin_pct, "Provision for adverse deviation on IBNR"),
        ("CAE %", cfg.cae_pct, "Claim adjustment expense on IBNR + margin"),
    ]
    for cat, r in v.categories.items():
        rows.append((f"Backlog gross-up ratio: {ABBREV[cat]}", r.restate_ratio, "Paid in valuation month / (paid + excess pended). 1.0 = no backlog"))
    names = {}
    for i, (k, val, note) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=k)
        c = ws.cell(row=i, column=2, value=val)
        c.fill = JUDGMENT
        if isinstance(val, float) and val < 1 and "%" in k:
            c.number_format = "0.0%"
        ws.cell(row=i, column=3, value=note)
        names[k] = f"Assumptions!$B${i}"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 70
    return names


def _triangle_sheet(wb: Workbook, name: str, r: CategoryResult) -> None:
    ws = wb.create_sheet(name)
    cum = r.triangle
    _header(ws, ["Incurred"] + [f"Lag {k}" for k in cum.columns])
    for i, (inc, row) in enumerate(cum.iterrows(), start=2):
        ws.cell(row=i, column=1, value=str(inc)).font = BOLD
        for j, val in enumerate(row, start=2):
            if pd.notna(val):
                c = ws.cell(row=i, column=j, value=round(float(val), 2))
                c.font = INPUT
                c.number_format = "#,##0"
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 10
    for j in range(2, cum.shape[1] + 2):
        ws.column_dimensions[col(j)].width = 13


def _ratio_sheet(wb: Workbook, name: str, tri: str, r: CategoryResult) -> None:
    ws = wb.create_sheet(name)
    cum = r.triangle
    n = cum.shape[1]
    _header(ws, ["Incurred"] + [f"{k}-{k + 1}" for k in range(n - 1)])
    for i, inc in enumerate(cum.index, start=2):
        ws.cell(row=i, column=1, value=str(inc)).font = BOLD
        for k in range(n - 1):
            if pd.isna(cum.iat[i - 2, k + 1]):
                continue
            a, b = _ref(tri, i, k + 2), _ref(tri, i, k + 3)
            c = ws.cell(row=i, column=k + 2, value=f'=IF({a}>0,{b}/{a},"")')
            c.number_format = "0.0000"
            if r.exclude_mask is not None and bool(r.exclude_mask.iat[i - 2, k]):
                c.fill = EXCLUDED
    ws.freeze_panes = "B2"


def _factor_sheet(wb: Workbook, name: str, tri: str, rat: str, r: CategoryResult, v: Valuation) -> int:
    """Selected factors by lag. Returns the first data row (lag 0)."""
    ws = wb.create_sheet(name)
    cfg = v.config
    cum = r.triangle
    n = cum.shape[1]
    row_of = {inc: i for i, inc in enumerate(cum.index, start=2)}
    volume = r.category == LARGE
    avg_label = "Volume-weighted avg (all history)" if volume else f"{cfg.cf_avg_months}-mo avg ex hi/lo"
    _header(ws, ["Lag", avg_label, "Override", "Selected link ratio", "Completion factor", "Months used"])
    overrides = cfg.factor_overrides.get(r.category, {})
    first, last = 2, n + 1
    for k in range(n):
        i = k + 2
        ws.cell(row=i, column=1, value=k)
        used = r.used_cells.get(k, []) if k < n - 1 else []
        if k >= cfg.max_lag or not used:
            avg = 1.0
        elif volume:
            num = ",".join(_ref(tri, row_of[m], k + 3) for m in used)
            den = ",".join(_ref(tri, row_of[m], k + 2) for m in used)
            avg = f"=SUM({num})/SUM({den})"
        else:
            cells = ",".join(_ref(rat, row_of[m], k + 2) for m in used)
            if cfg.cf_exclude_hi_lo and len(used) >= 4:
                avg = f"=(SUM({cells})-MAX({cells})-MIN({cells}))/(COUNT({cells})-2)"
            else:
                avg = f"=AVERAGE({cells})"
        ws.cell(row=i, column=2, value=avg).number_format = "0.0000"
        ov = ws.cell(row=i, column=3, value=overrides.get(k))
        ov.fill = JUDGMENT
        ov.number_format = "0.0000"
        ws.cell(row=i, column=4, value=f"=IF(ISNUMBER(C{i}),C{i},B{i})").number_format = "0.0000"
        ws.cell(row=i, column=5, value=f"=1/PRODUCT(D{i}:D{last})").number_format = "0.0%"
        ws.cell(row=i, column=6, value=", ".join(str(m) for m in used) if used else "complete")
    ws.freeze_panes = "B2"
    for c_, w in zip("ABCDEF", (6, 22, 11, 14, 14, 60)):
        ws.column_dimensions[c_].width = w
    return first


def _reserve_sheet(wb: Workbook, name: str, tri: str, fac: str, r: CategoryResult, names: dict[str, str]) -> tuple[int, int]:
    """Per incurred month reserve. Returns (first, last) data rows."""
    ws = wb.create_sheet(name)
    cum = r.triangle
    m = r.monthly
    large = r.category == LARGE
    _header(
        ws,
        ["Incurred", "Member months", "Age", "Paid to date", "Paid for estimation", "Completion factor", "CF ultimate",
         "Expected PMPM", "Expected claims", "B-F ultimate", "Selected ultimate", "IBNR", "Method"],
    )
    ratio = names[f"Backlog gross-up ratio: {ABBREV[r.category]}"]
    lags = names["PMPM / B-F months"]
    for i, (inc, row) in enumerate(m.iterrows(), start=2):
        age = int(row["age"])
        ws.cell(row=i, column=1, value=str(inc)).font = BOLD
        ws.cell(row=i, column=2, value=int(row["member_months"])).font = INPUT
        ws.cell(row=i, column=3, value=age)
        paid = _ref(tri, i, age + 2)
        prev = _ref(tri, i, age + 1) if age > 0 else "0"
        ws.cell(row=i, column=4, value=f"={paid}")
        ws.cell(row=i, column=5, value=f"={prev}+(D{i}-{prev})/{ratio}")
        ws.cell(row=i, column=6, value=f"=INDEX('{fac}'!$E:$E,C{i}+2)").number_format = "0.0%"
        ws.cell(row=i, column=7, value=f"=E{i}/F{i}")
        e = ws.cell(row=i, column=8, value=round(float(row["expected_claims"] / row["member_months"]), 6))
        e.font = INPUT
        e.number_format = "$#,##0.00"
        ws.cell(row=i, column=9, value=f"=H{i}*B{i}")
        ws.cell(row=i, column=10, value=f"=E{i}+(1-F{i})*I{i}")
        sel = f"=J{i}" if large else f"=IF(C{i}>={lags},G{i},J{i})"
        ws.cell(row=i, column=11, value=sel)
        ws.cell(row=i, column=12, value=f"=K{i}-D{i}")
        ws.cell(row=i, column=13, value="Large-claim B-F" if large else f'=IF(C{i}>={lags},"Completion factor","PMPM / B-F")')
        for c_ in (4, 5, 7, 9, 10, 11, 12):
            ws.cell(row=i, column=c_).number_format = "#,##0"
    first, last = 2, len(m) + 1
    t = last + 1
    ws.cell(row=t, column=1, value="Total").font = BOLD
    for c_ in (2, 4, 7, 9, 11, 12):
        cell = ws.cell(row=t, column=c_, value=f"=SUM({col(c_)}{first}:{col(c_)}{last})")
        cell.number_format = "#,##0"
        cell.font = BOLD
    ws.freeze_panes = "B2"
    widths = (10, 13, 6, 15, 15, 12, 15, 12, 15, 15, 15, 14, 18)
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[col(j)].width = w
    return first, last


def write(v: Valuation, path: Path) -> Path:
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    notes = [
        f"Health IBNR model: {v.valuation_month} valuation (live formulas)",
        "",
        "How to use",
        "1. Each category has Triangle (data), Ratios, Factors and Reserve sheets.",
        "2. Type a link ratio into a yellow Override cell on a Factors sheet. Completion factors, ultimates, IBNR and the Summary recalculate.",
        "3. Margin, CAE, the PMPM/B-F cut-off and backlog gross-up ratios are yellow cells on the Assumptions sheet.",
        "4. Grey ratios on the Ratios sheets are excluded from the averages (see the Selections log in the exhibits workbook).",
        "5. Expected PMPM (blue) comes from the Python trend fit: log-linear trend with calendar-month seasonality.",
        "",
        "Colour key: blue font = data input; yellow fill = judgment input; black = formula; grey fill = excluded ratio.",
        "Summary column 'Python engine' shows the model of record. 'Difference' should be zero until you change an input.",
    ]
    for i, t in enumerate(notes, start=1):
        c = readme.cell(row=i, column=1, value=t)
        if i in (1, 3):
            c.font = Font(bold=True, size=13 if i == 1 else 11)
    readme.column_dimensions["A"].width = 120

    summary = wb.create_sheet("Summary")
    names = _assumptions(wb, v)
    totals = {}
    for cat, r in v.categories.items():
        ab = ABBREV[cat]
        tri, rat, fac, res = f"{ab} Triangle", f"{ab} Ratios", f"{ab} Factors", f"{ab} Reserve"
        _triangle_sheet(wb, tri, r)
        _ratio_sheet(wb, rat, tri, r)
        _factor_sheet(wb, fac, tri, rat, r, v)
        _, last = _reserve_sheet(wb, res, tri, fac, r, names)
        totals[cat] = (f"'{res}'!D{last + 1}", f"'{res}'!L{last + 1}")

    _header(summary, ["Category", "Paid to date", "IBNR", "Margin", "CAE", "Total liability", "Python engine IBNR", "Difference"])
    py = v.summary()
    for i, (cat, (paid, ibnr)) in enumerate(totals.items(), start=2):
        summary.cell(row=i, column=1, value=v.label(cat))
        summary.cell(row=i, column=2, value=f"={paid}")
        summary.cell(row=i, column=3, value=f"={ibnr}")
        summary.cell(row=i, column=4, value=f"=C{i}*{names['Margin %']}")
        summary.cell(row=i, column=5, value=f"=(C{i}+D{i})*{names['CAE %']}")
        summary.cell(row=i, column=6, value=f"=C{i}+D{i}+E{i}")
        p = summary.cell(row=i, column=7, value=round(float(py.loc[v.label(cat), "ibnr_selected"]), 2))
        p.font = INPUT
        summary.cell(row=i, column=8, value=f"=C{i}-G{i}")
    t = len(totals) + 2
    summary.cell(row=t, column=1, value="Total").font = BOLD
    for c_ in range(2, 9):
        cell = summary.cell(row=t, column=c_, value=f"=SUM({col(c_)}2:{col(c_)}{t - 1})")
        cell.font = BOLD
    for row in summary.iter_rows(min_row=2, max_row=t, min_col=2, max_col=8):
        for c in row:
            c.number_format = "#,##0"
    summary.column_dimensions["A"].width = 24
    for j in range(2, 9):
        summary.column_dimensions[col(j)].width = 17

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
