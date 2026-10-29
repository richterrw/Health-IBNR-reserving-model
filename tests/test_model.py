import shutil
import subprocess
from dataclasses import replace

import numpy as np
import openpyxl
import pandas as pd
import pytest

from ibnr import excel_model, methods, synthetic, trend, triangle, valuation
from ibnr.config import ValuationConfig

VAL = pd.Period("2026-09", "M")
SELECTIONS = "selections.json"


@pytest.fixture(scope="module")
def data():
    membership, claims, truth, large, pended = synthetic.generate(valuation_month=str(VAL))
    book = triangle.Book(
        claims=triangle._periods(claims),
        members=pd.Series(membership["member_months"].values, index=pd.PeriodIndex(membership["incurred_month"], freq="M")),
        large_claims=triangle._periods(large),
        pended=triangle._periods(pended, ("paid_month",)),
    )
    true_ult = truth.groupby("incurred_month")["true_incurred"].sum()
    true_ult.index = pd.PeriodIndex(true_ult.index, freq="M")
    return book, truth, true_ult


@pytest.fixture(scope="module")
def cfg():
    return ValuationConfig().with_selections(SELECTIONS)


@pytest.fixture(scope="module")
def current(data, cfg):
    return valuation.run(data[0], VAL, cfg)


def true_ibnr(v, true_ult):
    tm = v.total_monthly
    return (true_ult.reindex(tm.index) - tm["paid"]).sum()


# ---------------------------------------------------------------- data

def test_lag_pattern_sums_to_one():
    p = synthetic.lag_pattern(2.0, 1.5, 36)
    assert p.sum() == pytest.approx(1.0)
    assert (p >= 0).all()


def test_triangle_shape_and_totals(data):
    claims = data[0].claims
    cum = triangle.build(claims, VAL, "Professional")
    n = cum.shape[0]
    assert cum.shape == (n, n)
    assert cum.iloc[-1].notna().sum() == 1  # valuation month: lag 0 only
    assert cum.iloc[0].notna().all()
    paid = triangle.latest_diagonal(cum).sum()
    assert paid == pytest.approx(claims.loc[claims["service_category"] == "Professional", "paid_amount"].sum())
    assert (cum.diff(axis=1).iloc[:, 1:].fillna(0) >= -1e-6).all().all()


def test_as_of_excludes_future_payments(data):
    earlier = triangle.build(data[0].claims, VAL - 3)
    assert earlier.index[-1] == VAL - 3
    assert triangle.ages(earlier).iloc[-1] == 0


def test_reconciliation_passes(data, current):
    recon = valuation.reconcile(data[0], current)
    assert (recon["result"] == "Pass").all(), recon


# ---------------------------------------------------------------- large claims

def test_large_claims_split_preserves_inpatient(data, current):
    book = data[0]
    ip_full = triangle.build(book.claims, VAL, "Inpatient")
    split = current.categories["Inpatient"].triangle + current.categories[valuation.LARGE].triangle
    np.testing.assert_allclose(split.fillna(0).to_numpy(), ip_full.fillna(0).to_numpy(), atol=0.05)


def test_large_claims_use_bf_everywhere(current):
    m = current.categories[valuation.LARGE].monthly
    assert (m["method"] == "Large-claim B-F").all()
    np.testing.assert_allclose(m["selected_ultimate"], m["bf_ultimate"])
    assert m["selected_ibnr"].sum() > 0


def test_without_large_claims_still_runs(data):
    v = valuation.run(replace(data[0], large_claims=None), VAL)
    assert valuation.LARGE not in v.categories
    assert v.label("Inpatient") == "Inpatient"


# ---------------------------------------------------------------- accuracy

def test_selected_ibnr_close_to_truth(data, current):
    tm = current.total_monthly
    assert tm["selected_ibnr"].sum() / true_ibnr(current, data[2]) - 1 == pytest.approx(0, abs=0.05)
    assert tm.loc[tm["age"] >= 12, "selected_ibnr"].sum() < 0.01 * tm["selected_ibnr"].sum()


def test_chain_ladder_overstates_when_claims_speed_up(current):
    s = current.summary().loc["Total"]
    assert s["ibnr_chain_ladder"] > s["ibnr_completion_factor"] > 0


def test_summary_loads(current):
    s = current.summary()
    cfg = current.config
    assert s.loc["Total", "ibnr_selected"] == pytest.approx(s.drop("Total")["ibnr_selected"].sum())
    t = s.loc["Total"]
    assert t["margin"] == pytest.approx(t["ibnr_selected"] * cfg.margin_pct)
    assert t["total_liability"] == pytest.approx(t["ibnr_selected"] + t["margin"] + t["cae"])
    assert t["chain_ladder_p75"] < t["chain_ladder_p95"]


def test_rollforward_identity(data, current, cfg):
    prior = valuation.run(data[0], VAL - 1, cfg)
    rf = valuation.rollforward(current, prior)
    for _, r in rf.iterrows():
        assert r["expected_remaining"] + r["development"] + r["new_month_ibnr"] == pytest.approx(r["current_ibnr"])
    assert rf.loc["Total", "prior_ibnr"] == pytest.approx(rf.drop("Total")["prior_ibnr"].sum())


# ---------------------------------------------------------------- judgment

def test_selections_load(cfg):
    assert cfg.has_judgment
    assert set(cfg.exclude_paid_months) == {"2026-06", "2026-07"}
    assert cfg.restate_paid_months == ("2026-06",)
    assert all(cfg.selection_notes.values())
    mech = cfg.mechanical()
    assert not mech.has_judgment and mech.margin_pct == cfg.margin_pct


def test_paid_month_mask_hits_the_diagonal():
    idx = pd.period_range("2026-01", "2026-06", freq="M")
    lr = pd.DataFrame(1.0, index=idx, columns=range(5))
    mask = methods.paid_month_mask(lr, ["2026-04"])
    # ratio (i, k) lands in paid month i + k + 1
    assert mask.loc[pd.Period("2026-01", "M"), 2] and mask.loc[pd.Period("2026-03", "M"), 0]
    assert mask.to_numpy().sum() == 3


def test_exclusions_skip_back_in_the_window(current):
    r = current.categories["Professional"]
    for k, months in r.used_cells.items():
        assert not any(r.exclude_mask.loc[m, k] for m in months)


def test_factor_override_applies(data, cfg):
    over = replace(cfg, factor_overrides={"Outpatient": {0: 3.0}})
    v = valuation.run(data[0], VAL, over)
    assert v.categories["Outpatient"].cf_link[0] == 3.0
    assert v.categories["Outpatient"].monthly["selected_ibnr"].sum() != pytest.approx(
        valuation.run(data[0], VAL, cfg).categories["Outpatient"].monthly["selected_ibnr"].sum()
    )


def test_backlog_restatement_fixes_the_backlog_close(data, cfg):
    book, truth, true_ult = data
    v_j = valuation.run(book, VAL - 3, cfg)  # June 2026 close, inside the backlog
    v_m = valuation.run(book, VAL - 3, cfg.mechanical())
    assert v_j.categories["Outpatient"].restate_ratio == pytest.approx(0.88, abs=0.005)
    err_j = abs(v_j.total_monthly["selected_ibnr"].sum() / true_ibnr(v_j, true_ult) - 1)
    err_m = abs(v_m.total_monthly["selected_ibnr"].sum() / true_ibnr(v_m, true_ult) - 1)
    assert err_j < err_m


def test_paid_month_diagnostic_flags_backlog(data):
    v = valuation.run(data[0], VAL)  # mechanical factors
    d = valuation.paid_month_diagnostic(v)
    assert d.loc["2026-06", "actual_to_expected"] < 0.95
    assert d.loc["2026-07", "actual_to_expected"] > 1.05


def test_hindsight_runs(data, current):
    book, truth, _ = data
    h = valuation.hindsight(book, current, range(4, 2, -1), truth)
    assert list(h.index) == ["2026-05", "2026-06"]
    assert {"true_ibnr", "ibnr_selected_error_pct", "ibnr_mechanical_error_pct"} <= set(h.columns)


# ---------------------------------------------------------------- trend

def test_trend_fit_recovers_known_trend():
    months = pd.period_range("2023-01", "2026-06", freq="M")
    season = pd.Series(synthetic.SEASONALITY, index=range(1, 13))
    pmpm = pd.Series(400 * 1.07 ** (np.arange(len(months)) / 12) * season.reindex(months.month).to_numpy(), index=months)
    fit = trend.fit(pmpm, pd.Series(True, index=months), fit_months=30, base_months=12)
    assert fit.annual_trend == pytest.approx(0.07, abs=1e-6)
    assert fit.r_squared == pytest.approx(1.0)
    np.testing.assert_allclose(fit.expected_pmpm(months[-6:]).to_numpy(), pmpm.iloc[-6:].to_numpy(), rtol=1e-9)


def test_trend_override(data):
    v = valuation.run(data[0], VAL, ValuationConfig(selected_annual_trend=0.0))
    assert all(r.trend.annual_trend == 0.0 for r in v.categories.values() if r.trend is not None)


# ---------------------------------------------------------------- live Excel model

def _recalc(path, out_dir):
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out_dir), str(path)],
                   check=True, capture_output=True, timeout=180)
    return openpyxl.load_workbook(out_dir / path.name, data_only=True)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_excel_model_ties_to_python(current, tmp_path):
    path = excel_model.write(current, tmp_path / "model.xlsx")
    wb = _recalc(path, tmp_path / "calc")
    rows = list(wb["Summary"].iter_rows(min_row=2, values_only=True))
    total = rows[-1]
    assert total[0] == "Total"
    assert total[2] == pytest.approx(current.summary().loc["Total", "ibnr_selected"], abs=5.0)
    assert all(abs(r[7]) < 5.0 for r in rows)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_excel_override_flows_through(data, current, cfg, tmp_path):
    path = excel_model.write(current, tmp_path / "model.xlsx")
    wb = openpyxl.load_workbook(path)
    wb["OP Factors"]["C2"] = 3.0  # lag 0 override
    wb.save(path)
    calc = _recalc(path, tmp_path / "calc")
    excel_op = next(r for r in calc["Summary"].iter_rows(min_row=2, values_only=True) if r[0] == "Outpatient")[2]
    py = valuation.run(data[0], VAL, replace(cfg, factor_overrides={"Outpatient": {0: 3.0}}))
    assert excel_op == pytest.approx(py.categories["Outpatient"].monthly["selected_ibnr"].sum(), abs=5.0)
